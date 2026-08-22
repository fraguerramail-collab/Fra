"""Esportazione della pianificazione mensile in un file Excel (.xlsx) con la
stessa impaginazione del foglio Google Sheets usato storicamente dal reparto
(titolo, gruppi di colonne colorati, righe weekend evidenziate, colonna
Assenze): l'utente lo scarica, lo perfeziona a mano se serve e lo salva in
PDF, senza dover toccare il foglio Google Drive originale a ogni mese.
"""

import calendar

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

WEEKDAY_ABBR = ["LUN", "MAR", "MER", "GIO", "VEN", "SAB", "DOM"]
SATURDAY = 5
SUNDAY = 6
MONDAY = 0
THURSDAY = 3

TITLE_FILL = "DD7E6B"
WEEKEND_FILL = "F9CB9C"
DAY_NUM_FILL = "F6B26B"

# (etichetta di gruppo o None, colonna iniziale, colonna finale, colore)
GROUPS = [
    ("Corsia", "D", "F", "D9EAD3"),
    (None, "G", "G", "D9EAD3"),  # Breast: stessa area/colore di Corsia, senza intestazione propria
    ("Giorno", "H", "J", "F4CCCC"),
    ("Notte", "K", "M", "C9DAF8"),
    ("Ambulatori", "N", "Q", "EAD1DC"),
    ("GOM", "R", "S", "C9DAF8"),
    ("MM", "T", "V", "FFE599"),
    ("ORB", "W", "Y", "B6D7A8"),
    ("CDP", "Z", "Z", "76A5AF"),
    ("ABB", "AA", "AA", "76A5AF"),
    ("Assenze", "AB", "AB", "D9D9D9"),
]

SUBHEADERS = {
    "D": "PO", "E": "Mod A", "F": "Mod B", "G": "Breast",
    "H": "GG", "I": "R1", "J": "R2/Urg",
    "K": "GN", "L": "R1", "M": "R2",
    "N": "Medic", "O": "Visite", "P": "Chir", "Q": "Proct",
    "R": "Sup", "S": "Inf",
    "T": "G", "U": "G", "V": "R",
    "W": "G", "X": "G", "Y": "R",
}

COLUMN_WIDTHS = {
    "A": 25.38, "B": 5.13, "C": 9.13, "D": 6.13, "E": 6.75, "F": 7.88, "G": 6.38,
    "H": 6.25, "I": 6.0, "J": 6.25, "K": 5.75, "L": 6.0, "M": 6.25, "N": 7.38,
    "O": 7.13, "P": 6.38, "Q": 6.75, "R": 7.13, "S": 7.25, "T": 5.38, "U": 4.88,
    "V": 13.0, "W": 4.38, "X": 5.38, "Y": 4.75, "Z": 6.25, "AA": 5.75, "AB": 38.63,
}

# colonna -> codice turno, per le colonne con un solo turno "semplice" (non
# dipendente dal giorno della settimana come GG/Visite, gestite a parte).
SIMPLE_COLUMN_SHIFT_CODES = {
    "D": "PO", "E": "MODA", "F": "MODB", "G": "BREAST",
    "I": "R1G", "J": "R2G", "K": "GN", "L": "R1N", "M": "R2N",
    "N": "MEDIC", "P": "CHIR", "Q": "PROC", "R": "GOMSUP", "S": "GOMINF",
    "T": "MAS", "U": "MAS_A", "V": "MAS_R", "W": "ORB", "X": "ORB_A", "Y": "ORB_R",
    "Z": "CDP", "AA": "ABB",
}

_thin = Side(style="thin")


def build_schedule_workbook(profile_name, year, month, month_name, types, employees, assignments, absences_by_day):
    """Costruisce il workbook per un mese. 'assignments' e' una lista di dict
    {employee_id, shift_type_id, day}; 'absences_by_day' e' un dict
    day -> testo gia' pronto (es. 'CAP, LOM - Ferie'), come nella pagina Assenze."""
    code_by_employee_id = {e.id: e.code for e in employees}
    shift_type_id_by_code = {st.code: st.id for st in types}

    assignees_by_slot = {}
    for a in assignments:
        assignees_by_slot.setdefault((a["shift_type_id"], a["day"]), []).append(
            code_by_employee_id.get(a["employee_id"], "?")
        )
    for key in assignees_by_slot:
        assignees_by_slot[key].sort()

    def codes_for(shift_code, day):
        st_id = shift_type_id_by_code.get(shift_code)
        if st_id is None:
            return ""
        return "/".join(assignees_by_slot.get((st_id, day), []))

    num_days = calendar.monthrange(year, month)[1]
    first_weekday = calendar.monthrange(year, month)[0]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = month_name[:3]

    for col, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[col].width = width
    ws.row_dimensions[1].height = 31.5
    ws.freeze_panes = "A4"

    last_col = "AB"
    ws.merge_cells(f"D1:{last_col}1")
    title_cell = ws["D1"]
    title_cell.value = f"{profile_name} - {month_name} {year}"
    title_cell.font = Font(name="Alata", size=14, bold=True)
    title_cell.fill = PatternFill("solid", fgColor=TITLE_FILL)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    title_cell.border = Border(top=_thin, left=_thin, right=_thin, bottom=_thin)

    for label, start, end, color in GROUPS:
        fill = PatternFill("solid", fgColor=color)
        if label is not None:
            if start != end:
                ws.merge_cells(f"{start}2:{end}2")
            header_cell = ws[f"{start}2"]
            header_cell.value = label
            header_cell.font = Font(name="Lexend", size=10, bold=True)
            header_cell.fill = fill
            header_cell.alignment = Alignment(horizontal="center", vertical="center")
            header_cell.border = Border(top=_thin)
        if start == end and start in ("Z", "AA", "AB"):
            ws.merge_cells(f"{start}2:{start}3")
            continue
        col_idx_start = openpyxl.utils.column_index_from_string(start)
        col_idx_end = openpyxl.utils.column_index_from_string(end)
        for col_idx in range(col_idx_start, col_idx_end + 1):
            col_letter = get_column_letter(col_idx)
            sub_cell = ws[f"{col_letter}3"]
            sub_cell.value = SUBHEADERS.get(col_letter, "")
            sub_cell.font = Font(name="Lexend", size=10, bold=True)
            sub_cell.fill = fill
            sub_cell.alignment = Alignment(horizontal="center", vertical="center")

    for day in range(1, num_days + 1):
        row = 3 + day
        weekday = (first_weekday + day - 1) % 7
        is_weekend = weekday in (SATURDAY, SUNDAY)

        day_cell = ws.cell(row=row, column=2, value=day)
        day_cell.font = Font(name="Lexend", size=10)
        day_cell.fill = PatternFill("solid", fgColor=DAY_NUM_FILL)
        day_cell.alignment = Alignment(horizontal="center")
        day_cell.border = Border(left=_thin)

        weekday_cell = ws.cell(row=row, column=3, value=WEEKDAY_ABBR[weekday])
        weekday_cell.font = Font(name="Lexend", size=10)
        weekday_cell.alignment = Alignment(horizontal="center", vertical="bottom")
        if is_weekend:
            weekday_cell.fill = PatternFill("solid", fgColor=WEEKEND_FILL)

        for label, start, end, color in GROUPS:
            col_idx_start = openpyxl.utils.column_index_from_string(start)
            col_idx_end = openpyxl.utils.column_index_from_string(end)
            for col_idx in range(col_idx_start, col_idx_end + 1):
                col_letter = get_column_letter(col_idx)
                cell = ws.cell(row=row, column=col_idx)
                cell.font = Font(name="Lexend", size=10)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if col_letter == "AB":
                    cell.value = absences_by_day.get(day, "")
                    cell.fill = PatternFill("solid", fgColor=color)
                    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                    continue
                cell.fill = PatternFill("solid", fgColor=WEEKEND_FILL if is_weekend else color)
                if col_letter == "H":
                    if weekday == SATURDAY:
                        parts = [p for p in (codes_for("GG_M", day), codes_for("GG_P", day)) if p]
                        cell.value = "/".join(parts)
                    else:
                        cell.value = codes_for("GG", day)
                elif col_letter == "O":
                    if weekday == MONDAY:
                        cell.value = codes_for("VISLUN", day)
                    elif weekday == THURSDAY:
                        cell.value = codes_for("VISGIO", day)
                    else:
                        cell.value = ""
                else:
                    shift_code = SIMPLE_COLUMN_SHIFT_CODES.get(col_letter)
                    cell.value = codes_for(shift_code, day) if shift_code else ""

    return wb
